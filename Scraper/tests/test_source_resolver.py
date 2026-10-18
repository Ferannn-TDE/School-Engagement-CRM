from contextlib import redirect_stdout
from io import BytesIO
from io import StringIO
import json
import re
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
import zipfile

import openpyxl

from source_resolver import serves_high_school
from models import School
from source_reader import SourceReader
from source_resolver import ColumnResolver, RainbowSource, SchoolSeedLookup, StateSourceResolver


class FakeResponse:
    def __init__(self, content=b"", url="https://example.test/data", content_type="application/octet-stream", data=None):
        self.content = content
        self.url = url
        self.headers = {"Content-Type": content_type}
        self.data = data
        self.text = content.decode("utf-8", errors="replace") if isinstance(content, bytes) else str(content)

    def raise_for_status(self):
        return None

    def json(self):
        return self.data if self.data is not None else json.loads(self.content)


class FakeSession:
    def __init__(self, responses):
        self.responses = responses
        self.headers = {}

    def get(self, url, params=None, **_kwargs):
        response = self.responses(url, params) if callable(self.responses) else self.responses[url]
        return response

    def post(self, url, **_kwargs):
        response = self.responses(url, None) if callable(self.responses) else self.responses[url]
        return response


def source_settings(url="https://example.test/data", **changes):
    settings = {
        "id": "test_schools",
        "enabled": True,
        "group": "public",
        "sector": "public",
        "url": url,
    }
    settings.update(changes)
    return settings


class ColumnResolverTests(unittest.TestCase):
    def test_padded_state_grade_codes_are_understood(self):
        self.assertTrue(serves_high_school("Central Academy", "", "090", "120"))
        self.assertTrue(serves_high_school("Central Academy", "0912"))
        self.assertTrue(serves_high_school("Central Academy", "KG12"))
        self.assertFalse(serves_high_school("Central Academy", "", "004", "080"))
        self.assertTrue(serves_high_school("Minneola High", ""))
        self.assertTrue(serves_high_school("Rolla JH/HS", ""))
        self.assertFalse(serves_high_school("Patton Jr High", ""))

    def test_aliases_resolve_and_unknown_fields_are_empty(self):
        rows = [{
            "Institution Name": "Central Academy",
            "Physical Street": "1 Main Street",
            "Unfamiliar Value": "ignored",
        }]
        columns = ColumnResolver(rows, {})
        self.assertEqual(columns.value(rows[0], "name"), "Central Academy")
        self.assertEqual(columns.value(rows[0], "address"), "1 Main Street")
        self.assertEqual(columns.value(rows[0], "administrator"), "")

    def test_short_school_column_does_not_become_state_or_status(self):
        rows = [{"School": "0001", "FacilityName": "Central High School"}]
        columns = ColumnResolver(rows, {"columns": {"name": "FacilityName"}})
        self.assertEqual(columns.value(rows[0], "state"), "")
        self.assertEqual(columns.value(rows[0], "status"), "")

    def test_license_state_does_not_become_the_address_state(self):
        rows = [{"School Name": "Central Academy", "Licensed in Other State": "Exempt"}]
        columns = ColumnResolver(rows, {})
        self.assertEqual(columns.value(rows[0], "state"), "")

    def test_headers_are_sampled_across_combined_sheets(self):
        rows = [{"District Name": f"District {number}"} for number in range(200)]
        rows.append({"SCHOOL_NAME": "Central High School", "High Grade": "12"})
        columns = ColumnResolver(rows, {})
        self.assertEqual(columns.find("name"), "SCHOOL_NAME")

    def test_shared_district_school_heading_can_be_a_school_name(self):
        rows = [{"District/School Name": "Central High School", "Grade Span": "9-12"}]
        columns = ColumnResolver(rows, {})
        self.assertEqual(columns.value(rows[0], "name"), "Central High School")

    def test_explicit_column_can_override_scope_guard(self):
        rows = [{"School Name": "Central Academy", "District Type": "Non-Public District"}]
        columns = ColumnResolver(rows, {"columns": {"school_type": "District Type"}})
        self.assertEqual(columns.value(rows[0], "school_type"), "Non-Public District")

    def test_excel_style_csv_text_is_unwrapped(self):
        rows = [{"School ID": '=\"000943\"'}]
        columns = ColumnResolver(rows, {})
        self.assertEqual(columns.value(rows[0], "state_id"), "000943")

    def test_configured_regex_extracts_part_of_a_field(self):
        rows = [{"School and Classification": "Central Academy Grades 9 through 12"}]
        columns = ColumnResolver(rows, {
            "columns": {"name": "School and Classification"},
            "extract": {"name": r"^(.*?)(?=\s+Grades?\s)"},
        })
        self.assertEqual(columns.value(rows[0], "name"), "Central Academy")


class RainbowSourceTests(unittest.TestCase):
    def test_cursor_paged_json_api_is_combined(self):
        first_url = "https://example.test/search?size=2"
        second_url = "https://example.test/search?size=2&cursor=next"
        responses = {
            first_url: FakeResponse(
                json.dumps({"items": [{"name": "A"}, {"name": "B"}], "cursor": "next"}).encode(),
                url=first_url,
                content_type="application/json",
            ),
            second_url: FakeResponse(
                json.dumps({"items": [{"name": "C"}], "cursor": ""}).encode(),
                url=second_url,
                content_type="application/json",
            ),
        }
        reader = SourceReader(
            source_settings(
                url="https://example.test/search",
                kind="json api",
                params={"size": 2},
                records_path="items",
                cursor_path="cursor",
            ),
            FakeSession(responses),
        )
        self.assertEqual([row["name"] for row in reader.rows()], ["A", "B", "C"])

    def test_seed_lookup_adds_website_profile_and_specific_query(self):
        url = "https://example.test/search?size=2&state=TX"
        content = json.dumps({
            "items": [{
                "name": "Central High School",
                "city": "Austin",
                "website": "central.example",
                "slug": "central-high-school-austin-tx",
            }],
            "cursor": "",
        }).encode()
        settings = {
            "url": "https://example.test/search",
            "kind": "json api",
            "params": {"size": 2},
            "records_path": "items",
            "cursor_path": "cursor",
            "state_parameter": "state",
            "columns": {"name": "name", "city": "city", "website": "website"},
            "profile_url": "https://directory.example/schools/{slug}",
            "query_url": "https://example.test/search?state={state}&search_term={name}",
        }
        school = School("TX:1", "Central High School", "TX", city="Austin")
        lookup = SchoolSeedLookup(
            "TX",
            settings,
            FakeSession({url: FakeResponse(content, url=url, content_type="application/json")}),
            {},
        )
        self.assertEqual(lookup.apply([school]), 1)
        self.assertEqual(school.website, "https://central.example/")
        self.assertIn("https://directory.example/schools/central-high-school-austin-tx", school.seeds)
        self.assertIn("search_term=Central+High+School", school.seeds[-1])

    def test_structured_url_and_joined_columns_are_resolved(self):
        content = json.dumps([{
            "Institution": "Central High School",
            "School ID": "10",
            "Grade Span": "9-12",
            "First": "Jane",
            "Last": "Doe",
            "Website": {"url": "https://central.example"},
        }]).encode()
        source = RainbowSource(
            "TX",
            source_settings(join_columns={"administrator": ["First", "Last"]}),
            FakeSession({"https://example.test/data": FakeResponse(content)}),
        )

        school = source.load()[0]
        self.assertEqual(school.administrator, "Jane Doe")
        self.assertEqual(school.website, "https://central.example/")

    def test_include_and_positive_row_rules_filter_before_mapping(self):
        content = (
            "School Name,School ID,Grade Span,Type,Grade 12\n"
            "Central High School,10,9-12,Public School,20\n"
            "Closed High School,11,9-12,Public LEA,20\n"
            "Empty High School,12,9-12,Public School,0\n"
        ).encode()
        source = RainbowSource(
            "TX",
            source_settings(
                include_rows=[{"column": "Type", "values": ["Public School"]}],
                positive_rows=[{"column": "Grade 12"}],
            ),
            FakeSession({"https://example.test/data": FakeResponse(content)}),
        )

        self.assertEqual([school.name for school in source.load()], ["Central High School"])

    def test_exclude_rule_accepts_a_regex(self):
        content = (
            "School Name,School ID,Grades Served\n"
            "Central High School,10,9-12\n"
            "Faith Christian School,11,9-12\n"
        ).encode()
        source = RainbowSource(
            "TX",
            source_settings(exclude_rows=[{
                "column": "School Name",
                "regex": r"\b(?:christian|catholic)\b",
            }]),
            FakeSession({"https://example.test/data": FakeResponse(content)}),
        )

        self.assertEqual([school.name for school in source.load()], ["Central High School"])

    def test_include_rule_can_use_a_regex(self):
        content = (
            "School Name,School ID,School Type\n"
            "Central Academy,10,Four-year High School\n"
            "West Academy,11,Elementary School\n"
        ).encode()
        source = RainbowSource(
            "TX",
            source_settings(
                high_school_only=True,
                include_rows=[{"column": "School Type", "regex": "high school|secondary"}],
            ),
            FakeSession({"https://example.test/data": FakeResponse(content)}),
        )
        self.assertEqual([school.name for school in source.load()], ["Central Academy"])

    def test_grade_12_enrollment_can_identify_a_high_school(self):
        content = (
            "School Name,School ID,Grade 12\n"
            "Central Academy,10,20\n"
            "West Academy,11,0\n"
        ).encode()
        source = RainbowSource(
            "TX",
            source_settings(grade_12_column="Grade 12"),
            FakeSession({"https://example.test/data": FakeResponse(content)}),
        )

        self.assertEqual([school.name for school in source.load()], ["Central Academy"])

    def test_grade_columns_build_the_served_range(self):
        content = (
            "School Name,School ID,K,1,9,10,11,12\n"
            "Central Academy,10,8,7,20,21,22,23\n"
        ).encode()
        source = RainbowSource(
            "TX",
            source_settings(
                grade_12_column="12",
                grade_columns=["K", "1", "9", "10", "11", "12"],
            ),
            FakeSession({"https://example.test/data": FakeResponse(content)}),
        )
        self.assertEqual(source.load()[0].grades, "K-12")

    def test_source_already_limited_to_high_schools_needs_no_grade_column(self):
        content = b"School Name,School ID\nCentral Academy,10\n"
        source = RainbowSource(
            "TX",
            source_settings(high_school_only=True),
            FakeSession({"https://example.test/data": FakeResponse(content)}),
        )
        self.assertEqual([school.name for school in source.load()], ["Central Academy"])

    def test_configured_private_flag_sets_school_type(self):
        content = (
            "School Name,School ID,Grade Span,Private Flag\n"
            "Central Academy,10,9-12,Yes\n"
        ).encode()
        source = RainbowSource(
            "TX",
            source_settings(
                sector="mixed",
                private_column="Private Flag",
                private_values=["Yes"],
            ),
            FakeSession({"https://example.test/data": FakeResponse(content)}),
        )

        self.assertEqual(source.load()[0].school_type, "Private high school")

    def test_school_and_district_columns_are_resolved_by_scope(self):
        content = (
            "County Name,District Number,District Name,District Web Page Address,School Number,"
            "School Name,School Street Address,School City,School Zip,"
            "School Web Page Address,School Principal,Grade Range,School Status\n"
            "CROSBY COUNTY,'001,CROSBYTON CISD,district.example,'010,CROSBYTON CISD PRE K-12,"
            "204 S HARRISON,CROSBYTON,'78701,school.example,MR DANIEL FITZGERALD,'EE-12,ACTIVE\n"
        ).encode()
        source = RainbowSource(
            "TX",
            source_settings(),
            FakeSession({"https://example.test/data": FakeResponse(content)}),
        )

        school = source.load()[0]
        self.assertEqual(school.facility_key, "TX:010")
        self.assertEqual(school.district_code, "001")
        self.assertEqual(school.name, "Crosbyton CISD Pre K-12")
        self.assertEqual(school.district_name, "Crosbyton CISD")
        self.assertEqual(school.county, "Crosby County")
        self.assertEqual(school.address, "204 S Harrison")
        self.assertEqual(school.city, "Crosbyton")
        self.assertEqual(school.zipcode, "78701")
        self.assertEqual(school.grades, "EE-12")
        self.assertEqual(school.administrator, "Mr. Daniel Fitzgerald")
        self.assertEqual(school.website, "https://school.example/")
        self.assertEqual(school.district_website, "https://district.example/")

    def test_combined_mailing_address_can_be_split(self):
        content = (
            "School Name,School ID,Grade Span,Mailing Address\n"
            'Central Academy,10,9-12,"4407 Payne Ave, Dayton, Ohio, 45414"\n'
        ).encode()
        source = RainbowSource(
            "OH",
            source_settings(split_address=True),
            FakeSession({"https://example.test/data": FakeResponse(content)}),
        )

        school = source.load()[0]
        self.assertEqual(school.address, "4407 Payne Ave")
        self.assertEqual(school.city, "Dayton")
        self.assertEqual(school.zipcode, "45414")

    def test_unknown_extension_falls_back_to_csv(self):
        content = (
            "Institution Name,District Code,School Code,Grade Span,Physical Street,City,Odd Column\n"
            "Central Academy,100,20,K-12,1 Main Street,Austin,anything\n"
            "West Middle School,100,30,6-8,2 Main Street,Austin,anything\n"
            "Central Adult School,100,40,9-12,3 Main Street,Austin,anything\n"
        ).encode()
        session = FakeSession({
            "https://example.test/data": FakeResponse(content),
        })
        source = RainbowSource("TX", source_settings(id_columns=["District Code", "School Code"]), session)
        schools = source.load()

        self.assertEqual(len(schools), 1)
        self.assertEqual(schools[0].facility_key, "TX:100-20")
        self.assertEqual(schools[0].address, "1 Main Street")
        self.assertEqual(schools[0].administrator, "")

    def test_configured_headings_help_find_the_real_header_row(self):
        content = (
            "Custom Institution,Custom Grade,Custom Identifier\n"
            "Central School,9-12,10\n"
            "West School,9-12,11\n"
        ).encode()
        source = RainbowSource(
            "TX",
            source_settings(columns={
                "name": "Custom Institution",
                "grades": "Custom Grade",
                "state_id": "Custom Identifier",
            }),
            FakeSession({"https://example.test/data": FakeResponse(content)}),
        )
        self.assertEqual(len(source.load()), 2)

    def test_nested_json_is_found_without_a_records_path(self):
        content = json.dumps({
            "result": {
                "features": [{
                    "attributes": {
                        "School Name": "North High School",
                        "School ID": "50",
                        "Grades Served": "9-12",
                        "State": "TX",
                    },
                }],
            },
        }).encode()
        source = RainbowSource(
            "TX",
            source_settings(),
            FakeSession({"https://example.test/data": FakeResponse(content)}),
        )
        self.assertEqual([school.name for school in source.load()], ["North High School"])

    def test_excel_finds_header_and_builds_composite_id(self):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.append(["Annual School Directory"])
        sheet.append([])
        sheet.append(["District Code", "Building Code", "School Name", "Grades Served", "Website"])
        sheet.append(["001", "020", "Central High School", "9-12", "central.example"])
        output = BytesIO()
        workbook.save(output)
        workbook.close()

        url = "https://example.test/schools.xlsx"
        source = RainbowSource(
            "TX",
            source_settings(url, id_columns=["District Code", "Building Code"]),
            FakeSession({url: FakeResponse(output.getvalue(), url, "application/octet-stream")}),
        )
        school = source.load()[0]
        self.assertEqual(school.facility_key, "TX:001-020")
        self.assertEqual(school.website, "https://central.example/")

    def test_excel_can_merge_two_header_rows(self):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.append(["School ID", "School Name", "Enrollment"])
        sheet.append([None, None, "Grade 12"])
        sheet.append(["10", "Central Academy", 20])
        output = BytesIO()
        workbook.save(output)
        workbook.close()

        url = "https://example.test/schools.xlsx"
        source = RainbowSource(
            "TX",
            source_settings(url, header_rows=[0, 1], grade_12_column="Grade 12"),
            FakeSession({url: FakeResponse(output.getvalue(), url)}),
        )

        self.assertEqual(source.load()[0].name, "Central Academy")

    def test_excel_with_empty_style_metadata_is_repaired(self):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.append(["School Name", "School ID", "Grades Served"])
        sheet.append(["Central High School", "10", "9-12"])
        valid = BytesIO()
        workbook.save(valid)
        workbook.close()

        broken = BytesIO()
        with zipfile.ZipFile(BytesIO(valid.getvalue())) as source:
            with zipfile.ZipFile(broken, "w") as target:
                for item in source.infolist():
                    content = source.read(item.filename)
                    if item.filename == "xl/styles.xml":
                        content = re.sub(
                            br"<fills[^>]*>.*?</fills>",
                            b"<fills><fill /></fills>",
                            content,
                        )
                    target.writestr(item, content)

        url = "https://example.test/schools.xlsx"
        source = RainbowSource(
            "TX",
            source_settings(url),
            FakeSession({url: FakeResponse(broken.getvalue(), url, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}),
        )
        self.assertEqual(source.load()[0].facility_key, "TX:10")

    def test_html_table_is_supported(self):
        content = b"""
            <table>
                <tr><th>School Name</th><th>High Grade</th><th>City</th></tr>
                <tr><td>Central High School</td><td>12</td><td>Austin</td></tr>
            </table>
        """
        url = "https://example.test/table"
        source = RainbowSource(
            "TX",
            source_settings(url),
            FakeSession({url: FakeResponse(content, url, "text/html")}),
        )
        self.assertEqual(source.load()[0].city, "Austin")

    def test_html_table_spans_keep_multiline_headers_aligned(self):
        content = b"""
            <table>
              <tr><th colspan='4'>Private Schools</th></tr>
              <tr><th rowspan='2'>Number</th><th colspan='2'>School</th><th rowspan='2'>Grades</th></tr>
              <tr><th>ID</th><th>Name</th></tr>
              <tr><td>1</td><td>10</td><td>Central Academy</td><td>9-12</td></tr>
            </table>
        """
        source = RainbowSource(
            "TX",
            source_settings(
                header_rows=[1, 2],
                columns={"state_id": "ID", "name": "Name", "grades": "Grades"},
            ),
            FakeSession({
                "https://example.test/data": FakeResponse(content, content_type="text/html"),
            }),
        )
        self.assertEqual(source.load()[0].facility_key, "TX:10")

    def test_html_css_selector_builds_records(self):
        content = b"""
            <ul>
              <li class='school'><a href='/profile/10'>Central Academy (10)</a></li>
              <li class='school'><a href='/profile/11'>West Academy (11)</a></li>
            </ul>
        """
        url = "https://example.test/directory"
        source = RainbowSource(
            "TX",
            source_settings(
                url,
                kind="html_select",
                row_selector="li.school a",
                high_school_only=True,
                fields={
                    "School Name": {"regex": r"^(.*?)\s*\("},
                    "School ID": {"attribute": "href", "regex": r"/(\d+)$"},
                    "Profile": {"attribute": "href", "resolve_url": True},
                },
            ),
            FakeSession({url: FakeResponse(content, url, "text/html")}),
        )
        schools = source.load()
        self.assertEqual([school.name for school in schools], ["Central Academy", "West Academy"])
        self.assertEqual(schools[0].alternate_seeds, ("https://example.test/profile/10",))

    def test_url_values_combine_repeated_directory_pages(self):
        template = "https://example.test/directory?district={value}"
        first = template.format(value="01")
        second = template.format(value="02")
        page = lambda name, number: (
            f"<a class='school' href='/profile/{number}'>{name}</a>".encode()
        )
        source = RainbowSource(
            "TX",
            source_settings(
                template,
                kind="html_select",
                url_values=["01", "02"],
                row_selector="a.school",
                high_school_only=True,
                fields={
                    "School Name": {},
                    "School ID": {"attribute": "href", "regex": r"/(\d+)$"},
                },
            ),
            FakeSession({
                first: FakeResponse(page("Central Academy", 10), first, "text/html"),
                second: FakeResponse(page("West Academy", 11), second, "text/html"),
            }),
        )
        self.assertEqual(len(source.load()), 2)

    def test_nested_layout_table_does_not_swallow_school_headers(self):
        content = b"""
            <table><tr><td>Directory
                <table>
                    <tr><th>School Name</th><th>High Grade</th></tr>
                    <tr><td>Central High School</td><td>12</td></tr>
                </table>
            </td></tr></table>
        """
        url = "https://example.test/table"
        source = RainbowSource(
            "TX",
            source_settings(url),
            FakeSession({url: FakeResponse(content, url, "text/html")}),
        )
        self.assertEqual(source.load()[0].name, "Central High School")

    def test_excel_content_type_with_html_fragment_is_supported(self):
        content = (
            b"<table><tr><th>School Name</th><th>School ID</th><th>Grade Span</th></tr>"
            b"<tr><td>Central High School</td><td>10</td><td>9-12</td></tr></table>"
        )
        source = RainbowSource(
            "TX",
            source_settings(),
            FakeSession({
                "https://example.test/data": FakeResponse(
                    content,
                    content_type="application/vnd.ms-excel",
                ),
            }),
        )

        self.assertEqual(source.load()[0].name, "Central High School")

    def test_xml_records_are_supported(self):
        content = b"""
            <schools><school><SchoolName>Central High School</SchoolName>
            <SchoolID>10</SchoolID><HighGrade>12</HighGrade></school></schools>
        """
        url = "https://example.test/schools.xml"
        source = RainbowSource(
            "TX",
            source_settings(url),
            FakeSession({url: FakeResponse(content, url, "application/xml")}),
        )
        self.assertEqual(source.load()[0].state_id, "10")

    def test_zip_archive_is_supported(self):
        output = BytesIO()
        with zipfile.ZipFile(output, "w") as archive:
            archive.writestr(
                "schools.csv",
                "School Name,School ID,Grades Served\nCentral High School,10,9-12\n",
            )
        url = "https://example.test/files.zip"
        source = RainbowSource(
            "TX",
            source_settings(url),
            FakeSession({url: FakeResponse(output.getvalue(), url, "application/zip")}),
        )
        self.assertEqual(source.load()[0].facility_key, "TX:10")

    def test_download_page_discovers_a_file(self):
        page = "https://example.test/directory"
        file_url = "https://example.test/current.csv"
        html = b'<a href="/current.csv">Current School Directory Download</a>'
        csv_file = b"School Name,School ID,Grades Served\nCentral High School,10,9-12\n"
        source = RainbowSource(
            "TX",
            source_settings(page),
            FakeSession({
                page: FakeResponse(html, page, "text/html"),
                file_url: FakeResponse(csv_file, file_url, "text/csv"),
            }),
        )
        self.assertEqual(source.load()[0].state_id, "10")

    def test_valid_page_table_is_used_before_download_links(self):
        page = "https://example.test/directory"
        html = b"""
            <a href='/partial.csv'>Download partial school data</a>
            <table>
              <tr><th>School Name</th><th>School ID</th><th>Grades</th></tr>
              <tr><td>Central High School</td><td>10</td><td>9-12</td></tr>
              <tr><td>West High School</td><td>11</td><td>9-12</td></tr>
            </table>
        """
        source = RainbowSource(
            "TX",
            source_settings(page),
            FakeSession({page: FakeResponse(html, page, "text/html")}),
        )

        self.assertEqual(len(source.load()), 2)

    def test_download_page_skips_a_district_only_file(self):
        page = "https://example.test/directory"
        district_url = "https://example.test/district.csv"
        school_url = "https://example.test/schools.csv"
        html = b"""
            <a href="/district.csv">Current District Directory Download</a>
            <a href="/schools.csv">Current School Directory Download</a>
        """
        source = RainbowSource(
            "TX",
            source_settings(page),
            FakeSession({
                page: FakeResponse(html, page, "text/html"),
                district_url: FakeResponse(
                    b"District Name,District ID\nExample District,1\n",
                    district_url,
                    "text/csv",
                ),
                school_url: FakeResponse(
                    b"School Name,School ID,Grades Served\nCentral High School,10,9-12\n",
                    school_url,
                    "text/csv",
                ),
            }),
        )
        self.assertEqual(source.load()[0].state_id, "10")

    def test_arcgis_pages_are_combined(self):
        url = "https://example.test/MapServer/0"

        def responses(request_url, params):
            self.assertEqual(request_url, url + "/query")
            offset = params["resultOffset"]
            features = []
            if offset == 0:
                features = [{"attributes": {
                    "School Name": "Central High School",
                    "School ID": "10",
                    "Grades Served": "9-12",
                }}]
            return FakeResponse(data={"features": features})

        settings = source_settings(url, kind="arcgis", page_size=1)
        source = RainbowSource("TX", settings, FakeSession(responses))
        self.assertEqual(source.load()[0].state_id, "10")

    def test_unfamiliar_url_field_becomes_an_inferred_seed(self):
        content = (
            "School Name,School ID,Grades Served,Portal Destination\n"
            "Central High School,10,9-12,central.example\n"
        ).encode()
        source = RainbowSource(
            "TX",
            source_settings(),
            FakeSession({"https://example.test/data": FakeResponse(content)}),
        )
        school = source.load()[0]
        self.assertEqual(school.website, "")
        self.assertEqual(school.alternate_seeds, ("https://central.example/",))

    def test_powerapps_records_include_websites_and_profile_seeds(self):
        landing_url = "https://directory.example.test/results"
        token_url = "https://directory.example.test/_services/auth/portalusertoken"
        api_url = "https://directory.example.test/_api/accounts"
        settings = source_settings(
            landing_url,
            kind="powerapps",
            fields=["accountid", "school_name", "grades", "websiteurl"],
            filters=[{"field": "statecode", "operator": "eq", "value": "0"}],
            profile_url="https://directory.example.test/profile?id={accountid}",
            columns={
                "name": "school_name",
                "grades": "grades",
                "website": "websiteurl",
            },
        )
        session = FakeSession({
            landing_url: FakeResponse(b"<html></html>", landing_url, "text/html"),
            token_url: FakeResponse(b"token", token_url, "text/plain"),
            api_url: FakeResponse(data={"value": [{
                "accountid": "abc",
                "school_name": "Central High School",
                "grades": "9, 10, 11, 12",
                "websiteurl": "central.example",
            }]}),
        })
        school = RainbowSource("TN", settings, session).load()[0]
        self.assertEqual(school.website, "https://central.example/")
        self.assertEqual(
            school.alternate_seeds,
            ("https://directory.example.test/profile?id=abc",),
        )

    def test_form_export_is_parsed(self):
        landing_url = "https://directory.example.test/search"
        export_url = "https://directory.example.test/export"
        page = b'<form action="/export"><input type="hidden" name="token" value="abc"></form>'
        export = b"School Name,School ID,Grades Served\nCentral High School,10,9-12\n"
        source = RainbowSource(
            "TX",
            source_settings(
                landing_url,
                kind="form",
                posts=[{"url": "/export", "values": {"button": "Download"}}],
            ),
            FakeSession({
                landing_url: FakeResponse(page, landing_url, "text/html"),
                export_url: FakeResponse(export, export_url, "text/csv"),
            }),
        )
        self.assertEqual(source.load()[0].facility_key, "TX:10")

    def test_form_can_export_each_select_option(self):
        landing = b"""
            <form action='/export'>
              <input type='hidden' name='token' value='abc'>
              <select name='region'>
                <option value='1'>North</option>
                <option value='2'>South</option>
              </select>
            </form>
        """

        class OptionSession(FakeSession):
            def post(self, url, data=None, **_kwargs):
                school = "North" if data["region"] == "1" else "South"
                return FakeResponse(
                    f"School Name,School ID,Grade Span\n{school} High School,{data['region']},9-12\n".encode(),
                    url,
                    "text/csv",
                )

        source = RainbowSource(
            "TX",
            source_settings(
                "https://example.test/form",
                kind="form",
                posts=[{"for_each_option": "region"}],
            ),
            OptionSession({
                "https://example.test/form": FakeResponse(
                    landing,
                    "https://example.test/form",
                    "text/html",
                ),
            }),
        )

        self.assertEqual([school.name for school in source.load()], [
            "North High School",
            "South High School",
        ])

    def test_form_can_follow_postback_pages(self):
        landing_url = "https://example.test/form"
        landing = b"<form><input type='hidden' name='token' value='first'></form>"
        first = b"""
            <form><input type='hidden' name='token' value='next'>
            <table><tr><th>School Name</th><th>Grade Span</th></tr>
            <tr><td>North High School</td><td>9-12</td></tr></table>
            <div class='pages'>
              <a class='current' href="javascript:__doPostBack('page1','')">1</a>
              <a href="javascript:__doPostBack('page2','')">2</a>
            </div></form>
        """
        second = b"""
            <table><tr><th>School Name</th><th>Grade Span</th></tr>
            <tr><td>South High School</td><td>9-12</td></tr></table>
        """

        class PagerSession(FakeSession):
            def post(self, url, data=None, **_kwargs):
                if data.get("__EVENTTARGET") == "page2":
                    self.test.assertEqual(data["school_type"], "public")
                    return FakeResponse(second, url, "text/html")
                return FakeResponse(first, url, "text/html")

        session = PagerSession({
            landing_url: FakeResponse(landing, landing_url, "text/html"),
        })
        session.test = self
        source = RainbowSource(
            "TX",
            source_settings(
                landing_url,
                kind="form",
                posts=[{"values": {"search": "Search"}}],
                postback_selector=".pages a",
                postback_values={"school_type": "public"},
            ),
            session,
        )

        self.assertEqual([school.name for school in source.load()], [
            "North High School",
            "South High School",
        ])


class StateSourceResolverTests(unittest.TestCase):
    def test_nested_source_template_inherits_group_and_state_parameter(self):
        data = {
            "directory": {
                "url": "https://example.test/search",
                "kind": "json api",
                "state_parameter": "state",
            },
            "private_directory": {
                "template": "directory",
                "group": "private",
                "sector": "private",
            },
            "states": {
                "TX": {
                    "enabled": True,
                    "sources": [{"id": "tx_private", "template": "private_directory"}],
                },
            },
        }
        resolver = StateSourceResolver(data=data, session=FakeSession({}))
        source = resolver.source_groups("TX", data["states"]["TX"])["private"][0]
        self.assertEqual(source["url"], "https://example.test/search")
        self.assertEqual(source["sector"], "private")
        self.assertEqual(source["params"], {"state": "TX"})

    def test_private_source_cannot_duplicate_the_public_roster(self):
        public_url = "https://example.test/public.csv"
        private_url = "https://example.test/private.csv"
        data = {
            "states": {
                "TX": {
                    "enabled": True,
                    "sources": [
                        source_settings(public_url, id="public"),
                        source_settings(
                            private_url,
                            id="private",
                            group="private",
                            sector="private",
                        ),
                    ],
                },
            },
        }
        session = FakeSession({
            public_url: FakeResponse(
                b"School Name,School ID,Grades Served\nCentral High School,1,9-12\n",
                public_url,
                "text/csv",
            ),
            private_url: FakeResponse(
                b"School Name,School ID,Grades Served\n"
                b"Central High School,2,9-12\nFaith Academy,3,9-12\n",
                private_url,
                "text/csv",
            ),
        })
        with tempfile.TemporaryDirectory() as folder:
            schools = StateSourceResolver(data=data, session=session).load(folder)

        self.assertEqual([school.name for school in schools], [
            "Central High School",
            "Faith Academy",
        ])
        self.assertEqual(schools[1].school_type, "Private high school")

    def test_enabled_states_and_environment_override(self):
        data = {
            "states": {
                "TX": {"enabled": True, "sources": [source_settings()]},
                "IA": {"enabled": False, "sources": [source_settings()]},
            },
        }
        resolver = StateSourceResolver(data=data, session=FakeSession({}), states=None)
        self.assertEqual(resolver.states, ["TX"])
        with patch.dict("os.environ", {"SCHOOL_STATES": "ia,TX"}):
            resolver = StateSourceResolver(data=data, session=FakeSession({}))
        self.assertEqual(resolver.states, ["IA", "TX"])

    def test_fallback_and_minimal_health_file(self):
        bad_url = "https://example.test/bad"
        good_url = "https://example.test/good.csv"
        data = {
            "states": {
                "TX": {
                    "enabled": True,
                    "sources": [
                        source_settings(bad_url, id="bad", priority=1),
                        source_settings(good_url, id="good", priority=2),
                    ],
                },
            },
        }
        session = FakeSession({
            bad_url: FakeResponse(b"not a school source", bad_url),
            good_url: FakeResponse(
                b"School Name,School ID,Grades Served\nCentral High School,10,9-12\n",
                good_url,
                "text/csv",
            ),
        })
        output = StringIO()
        with redirect_stdout(output), tempfile.TemporaryDirectory() as folder:
            schools = StateSourceResolver(data=data, session=session).load(folder)
            health = json.loads((Path(folder) / "source_health.json").read_text())

        self.assertEqual([school.state_id for school in schools], ["10"])
        self.assertEqual(output.getvalue().strip(), "TX 1 schools added")
        self.assertEqual(health["TX"]["bad"]["success"], False)
        self.assertEqual(health["TX"]["good"]["success"], True)
        self.assertEqual(set(health["TX"]["good"]), {"success", "checked_at"})

    def test_larger_source_replaces_a_partial_source(self):
        partial_url = "https://example.test/partial.csv"
        complete_url = "https://example.test/complete.csv"
        partial = b"School Name,School ID,Grades Served\nFirst High School,1,9-12\n"
        complete = (
            b"School Name,School ID,Grades Served\n"
            b"First High School,1,9-12\nSecond High School,2,9-12\n"
        )
        data = {
            "states": {
                "TX": {
                    "enabled": True,
                    "sources": [
                        source_settings(partial_url, id="partial", priority=1),
                        source_settings(complete_url, id="complete", priority=2),
                    ],
                },
            },
        }
        session = FakeSession({
            partial_url: FakeResponse(partial, partial_url, "text/csv"),
            complete_url: FakeResponse(complete, complete_url, "text/csv"),
        })
        with tempfile.TemporaryDirectory() as folder:
            schools = StateSourceResolver(data=data, session=session).load(folder)

        self.assertEqual([school.state_id for school in schools], ["1", "2"])

    def test_all_selected_states_are_checked_before_failure(self):
        bad_url = "https://example.test/bad.csv"
        good_url = "https://example.test/good.csv"
        data = {
            "states": {
                "TX": {
                    "enabled": True,
                    "sources": [source_settings(bad_url, id="bad")],
                },
                "IA": {
                    "enabled": True,
                    "sources": [source_settings(good_url, id="good")],
                },
            },
        }
        session = FakeSession({
            bad_url: FakeResponse(b"not a school source", bad_url),
            good_url: FakeResponse(
                b"School Name,School ID,Grades Served\nCentral High School,10,9-12\n",
                good_url,
                "text/csv",
            ),
        })
        output = StringIO()
        with redirect_stdout(output), tempfile.TemporaryDirectory() as folder:
            with self.assertRaisesRegex(RuntimeError, "Sources failed for 1 states: TX"):
                StateSourceResolver(data=data, session=session).load(folder)
            health = json.loads((Path(folder) / "source_health.json").read_text())
            self.assertFalse((Path(folder) / "normalized_schools.json").exists())

        self.assertEqual(output.getvalue().splitlines(), [
            "TX source failed",
            "IA 1 schools added",
        ])
        self.assertFalse(health["TX"]["bad"]["success"])
        self.assertTrue(health["IA"]["good"]["success"])


if __name__ == "__main__":
    unittest.main()
